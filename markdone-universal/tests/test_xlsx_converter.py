#!/usr/bin/env python3

"""
Comprehensive pytest test suite for xlsx_converter.py

Run with:
    python -m pytest tests/test_xlsx_converter.py -v
"""

import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import openpyxl
import pytest


# ── Helpers ───────────────────────────────────────────────────────────────────

def create_xlsx_fixture(tmp_path: Path, filename: str, sheets_data: dict[str, list[list[Any]]]) -> Path:
    """
    Create an .xlsx file with the given sheets and data.
    
    Args:
        tmp_path: pytest tmp_path fixture
        filename: name of the xlsx file to create
        sheets_data: dict mapping sheet names to 2D lists of cell values
                     First row is treated as headers
    
    Returns:
        Path to the created .xlsx file
    """
    xlsx_path = tmp_path / filename
    wb = openpyxl.Workbook()
    
    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    
    for sheet_name, data in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name)
        for row in data:
            ws.append(row)
    
    wb.save(str(xlsx_path))
    wb.close()
    return xlsx_path


def run_converter(
    input_path: Path,
    output_path: Path,
    target: str = "vault_md",
    yaml_frontmatter: str = "true",
) -> tuple[int, str, str]:
    """
    Run xlsx_converter.py and return (exit_code, stdout, stderr).
    """
    script_path = Path(__file__).parent.parent / "services" / "xlsx_converter.py"
    
    cmd = [
        sys.executable,
        str(script_path),
        "--input", str(input_path),
        "--output", str(output_path),
        "--target", target,
        "--yaml-frontmatter", yaml_frontmatter,
    ]
    
    result = subprocess.run(
        cmd,
        capture_output=True,
        text=True,
    )
    
    return result.returncode, result.stdout, result.stderr


def parse_json_output(stdout: str) -> dict:
    """Parse the single JSON line from stdout."""
    lines = [line.strip() for line in stdout.strip().split("\n") if line.strip()]
    if not lines:
        raise ValueError("No JSON output found in stdout")
    return json.loads(lines[0])


# ══════════════════════════════════════════════════════════════════════════════
# CONTRACT TESTS — Verify exact JSON payload fields Deno depends on
# ══════════════════════════════════════════════════════════════════════════════

def test_contract_vault_md_success_payload(tmp_path: Path):
    """Verify vault_md success payload contains all required fields."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"Sheet1": [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]]},
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path, target="vault_md")
    
    assert exit_code == 0, f"Expected exit code 0, got {exit_code}"
    
    payload = parse_json_output(stdout)
    
    # Verify all required fields are present
    assert "status" in payload
    assert "outputPath" in payload
    assert "sourceType" in payload
    assert "sheetsProcessed" in payload
    assert "rowsProcessed" in payload
    assert "durationMs" in payload
    
    # Verify field values
    assert payload["status"] == "completed"
    assert payload["outputPath"] == str(output_path)
    assert payload["sourceType"] == "xlsx"
    assert payload["sheetsProcessed"] == 1
    assert payload["rowsProcessed"] == 2
    assert isinstance(payload["durationMs"], int)
    assert payload["durationMs"] >= 0


def test_contract_vault_json_success_payload(tmp_path: Path):
    """Verify vault_json success payload contains all required fields."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"Sheet1": [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]]},
    )
    output_path = tmp_path / "output.json"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path, target="vault_json")
    
    assert exit_code == 0, f"Expected exit code 0, got {exit_code}"
    
    payload = parse_json_output(stdout)
    
    # Verify all required fields are present
    assert "status" in payload
    assert "outputPath" in payload
    assert "sourceType" in payload
    assert "sheetsProcessed" in payload
    assert "rowsProcessed" in payload
    assert "durationMs" in payload
    
    # Verify field values
    assert payload["status"] == "completed"
    assert payload["outputPath"] == str(output_path)
    assert payload["sourceType"] == "xlsx"
    assert payload["sheetsProcessed"] == 1
    assert payload["rowsProcessed"] == 2
    assert isinstance(payload["durationMs"], int)
    assert payload["durationMs"] >= 0


# ══════════════════════════════════════════════════════════════════════════════
# MARKDOWN OUTPUT TESTS
# ══════════════════════════════════════════════════════════════════════════════

def test_markdown_yaml_frontmatter_enabled(tmp_path: Path):
    """Verify YAML frontmatter is present when enabled."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"Sheet1": [["Name", "Age"], ["Alice", "30"]]},
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="true"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify frontmatter structure
    assert content.startswith("---\n")
    assert "title:" in content
    assert "source:" in content
    assert "lang: en" in content
    assert "converted_at:" in content
    assert "sheets:" in content
    assert "sheet_types:" in content
    
    # Verify frontmatter ends properly
    lines = content.split("\n")
    frontmatter_end = None
    for i, line in enumerate(lines[1:], start=1):
        if line == "---":
            frontmatter_end = i
            break
    
    assert frontmatter_end is not None, "Frontmatter closing --- not found"


def test_markdown_yaml_frontmatter_disabled(tmp_path: Path):
    """Verify YAML frontmatter is absent when disabled."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"Sheet1": [["Name", "Age"], ["Alice", "30"]]},
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify no frontmatter
    assert not content.startswith("---\n")
    assert "lang: en" not in content.split("\n")[0:10]  # Check first 10 lines
    
    # Should start with main heading
    assert content.startswith("# test\n")


def test_markdown_lang_en_in_frontmatter(tmp_path: Path):
    """Verify lang: en is present in frontmatter."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"Sheet1": [["Name", "Age"], ["Alice", "30"]]},
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="true"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    assert "lang: en\n" in content


def test_markdown_sheet_name_becomes_heading(tmp_path: Path):
    """Verify sheet name becomes ## heading."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"MySheet": [["Name", "Age"], ["Alice", "30"]]},
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    assert "## MySheet\n" in content


def test_markdown_table_structure(tmp_path: Path):
    """Verify Markdown table structure is correct."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "Sheet1": [
                ["Name", "Age", "City"],
                ["Alice", "30", "NYC"],
                ["Bob", "25", "LA"],
            ]
        },
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify table structure
    assert "| Name | Age | City |" in content
    assert "| --- | --- | --- |" in content
    assert "| Alice | 30 | NYC |" in content
    assert "| Bob | 25 | LA |" in content


def test_markdown_multiple_sheets(tmp_path: Path):
    """Verify multiple sheets all appear in output."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "Sheet1": [["Name"], ["Alice"]],
            "Sheet2": [["Age"], ["30"]],
            "Sheet3": [["City"], ["NYC"]],
        },
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify all sheets are present
    assert "## Sheet1\n" in content
    assert "## Sheet2\n" in content
    assert "## Sheet3\n" in content


# ══════════════════════════════════════════════════════════════════════════════
# JSON OUTPUT TESTS
# ══════════════════════════════════════════════════════════════════════════════

def test_json_output_structure(tmp_path: Path):
    """Verify JSON output has correct structure."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "Sheet1": [
                ["Name", "Age"],
                ["Alice", "30"],
                ["Bob", "25"],
            ]
        },
    )
    output_path = tmp_path / "output.json"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path, target="vault_json")
    
    assert exit_code == 0
    
    with open(output_path) as f:
        data = json.load(f)
    
    # Verify top-level structure
    assert "source" in data
    assert "lang" in data
    assert "converted_at" in data
    assert "sheets" in data
    
    # Verify source and lang
    assert data["source"] == "test.xlsx"
    assert data["lang"] == "en"
    
    # Verify sheets structure
    assert "Sheet1" in data["sheets"]
    sheet = data["sheets"]["Sheet1"]
    
    assert "type" in sheet
    assert "columns" in sheet
    assert "row_count" in sheet
    assert "rows" in sheet
    
    # Verify columns
    assert sheet["columns"] == ["Name", "Age"]
    
    # Verify row count
    assert sheet["row_count"] == 2
    
    # Verify rows
    assert len(sheet["rows"]) == 2
    assert sheet["rows"][0] == {"Name": "Alice", "Age": "30"}
    assert sheet["rows"][1] == {"Name": "Bob", "Age": "25"}


def test_json_lang_en_present(tmp_path: Path):
    """Verify lang: en field is present in JSON output."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"Sheet1": [["Name"], ["Alice"]]},
    )
    output_path = tmp_path / "output.json"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path, target="vault_json")
    
    assert exit_code == 0
    
    with open(output_path) as f:
        data = json.load(f)
    
    assert data["lang"] == "en"


# ══════════════════════════════════════════════════════════════════════════════
# COMPLIANCE DETECTION TESTS
# ══════════════════════════════════════════════════════════════════════════════

def test_compliance_detection_positive(tmp_path: Path):
    """Verify sheet with compliance keywords is flagged as compliance matrix."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "ComplianceSheet": [
                ["Requirement", "Status", "Vendor", "Response"],
                ["REQ-001", "Met", "Acme Corp", "Compliant"],
                ["REQ-002", "Partially", "Beta Inc", "In Progress"],
            ]
        },
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify compliance badge is present
    assert "## ComplianceSheet _(compliance matrix)_" in content


def test_compliance_detection_negative(tmp_path: Path):
    """Verify generic sheet is not flagged as compliance matrix."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "GenericSheet": [
                ["Name", "Age", "City"],
                ["Alice", "30", "NYC"],
                ["Bob", "25", "LA"],
            ]
        },
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify no compliance badge
    assert "## GenericSheet\n" in content
    assert "_(compliance matrix)_" not in content


def test_compliance_detection_in_json(tmp_path: Path):
    """Verify compliance detection works in JSON output."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "ComplianceSheet": [
                ["Requirement", "Compliance", "Status"],
                ["REQ-001", "Yes", "Met"],
            ],
            "GenericSheet": [
                ["Name", "Age"],
                ["Alice", "30"],
            ],
        },
    )
    output_path = tmp_path / "output.json"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path, target="vault_json")
    
    assert exit_code == 0
    
    with open(output_path) as f:
        data = json.load(f)
    
    assert data["sheets"]["ComplianceSheet"]["type"] == "compliance"
    assert data["sheets"]["GenericSheet"]["type"] == "generic"


# ══════════════════════════════════════════════════════════════════════════════
# ERROR HANDLING TESTS
# ══════════════════════════════════════════════════════════════════════════════

def test_error_missing_input_file(tmp_path: Path):
    """Verify missing input file returns exit code 1 with INPUT_NOT_FOUND."""
    input_path = tmp_path / "nonexistent.xlsx"
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(input_path, output_path)
    
    assert exit_code == 1
    
    # Error should be in stderr
    assert "INPUT_NOT_FOUND" in stderr
    assert str(input_path) in stderr


def test_error_unsupported_extension(tmp_path: Path):
    """Verify unsupported extension returns exit code 1 with UNSUPPORTED_FILE_TYPE."""
    input_path = tmp_path / "test.txt"
    input_path.write_text("not an xlsx file")
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(input_path, output_path)
    
    assert exit_code == 1
    
    # Error should be in stderr
    assert "UNSUPPORTED_FILE_TYPE" in stderr


def test_error_empty_workbook(tmp_path: Path):
    """Verify empty workbook returns exit code 2 with EXTRACTION_FAILED."""
    # Create a workbook with an empty sheet (no data)
    xlsx_path = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    # Keep the default sheet but don't add any data to it
    # This will result in a workbook with no readable content
    wb.save(str(xlsx_path))
    wb.close()
    
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path)
    
    assert exit_code == 2
    
    # Error should be in stderr
    assert "EXTRACTION_FAILED" in stderr


# ══════════════════════════════════════════════════════════════════════════════
# EDGE CASES
# ══════════════════════════════════════════════════════════════════════════════

def test_edge_case_pipe_character_escaping(tmp_path: Path):
    """Verify pipe characters | in cells are escaped as \\|."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "Sheet1": [
                ["Name", "Description"],
                ["Alice", "Uses | pipe character"],
                ["Bob", "Multiple | pipes | here"],
            ]
        },
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify pipes are escaped
    assert "Uses \\| pipe character" in content
    assert "Multiple \\| pipes \\| here" in content


def test_edge_case_cell_truncation(tmp_path: Path):
    """Verify cells over 500 characters are truncated with …."""
    long_text = "A" * 600
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "Sheet1": [
                ["Name", "Description"],
                ["Alice", long_text],
            ]
        },
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify truncation
    assert ("A" * 500 + "…") in content
    assert ("A" * 501) not in content


def test_edge_case_row_cap(tmp_path: Path):
    """Verify sheets over 2000 rows are capped."""
    # Create a sheet with 2500 rows
    rows = [["Index", "Value"]]
    for i in range(2500):
        rows.append([str(i), f"Value{i}"])
    
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"Sheet1": rows},
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    
    # Verify row count in payload
    payload = parse_json_output(stdout)
    assert payload["rowsProcessed"] == 2000
    
    # Verify content doesn't include rows beyond 2000
    content = output_path.read_text()
    assert "Value1999" in content
    assert "Value2000" not in content
    assert "Value2499" not in content


def test_edge_case_row_cap_json(tmp_path: Path):
    """Verify row cap works in JSON output."""
    # Create a sheet with 2500 rows
    rows = [["Index", "Value"]]
    for i in range(2500):
        rows.append([str(i), f"Value{i}"])
    
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {"Sheet1": rows},
    )
    output_path = tmp_path / "output.json"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path, target="vault_json")
    
    assert exit_code == 0
    
    with open(output_path) as f:
        data = json.load(f)
    
    # Verify row count is capped at 2000
    assert data["sheets"]["Sheet1"]["row_count"] == 2000
    assert len(data["sheets"]["Sheet1"]["rows"]) == 2000


def test_edge_case_newlines_in_cells(tmp_path: Path):
    """Verify newlines in cells are replaced with spaces."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "Sheet1": [
                ["Name", "Description"],
                ["Alice", "Line 1\nLine 2\nLine 3"],
            ]
        },
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify newlines are replaced with spaces
    assert "Line 1 Line 2 Line 3" in content
    assert "Line 1\nLine 2" not in content


def test_edge_case_empty_cells(tmp_path: Path):
    """Verify empty cells are handled correctly."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "Sheet1": [
                ["Name", "Age", "City"],
                ["Alice", "", "NYC"],
                ["", "25", ""],
            ]
        },
    )
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(
        xlsx_path, output_path, target="vault_md", yaml_frontmatter="false"
    )
    
    assert exit_code == 0
    content = output_path.read_text()
    
    # Verify empty cells are represented as empty strings in table
    assert "| Alice |  | NYC |" in content
    assert "|  | 25 |  |" in content


# ══════════════════════════════════════════════════════════════════════════════
# ADDITIONAL INTEGRATION TESTS
# ══════════════════════════════════════════════════════════════════════════════

def test_multiple_sheets_with_mixed_types(tmp_path: Path):
    """Test workbook with both compliance and generic sheets."""
    xlsx_path = create_xlsx_fixture(
        tmp_path,
        "test.xlsx",
        {
            "Compliance": [
                ["Requirement", "Status", "Vendor"],
                ["REQ-001", "Met", "Acme"],
            ],
            "Data": [
                ["Name", "Age"],
                ["Alice", "30"],
            ],
            "Audit": [
                ["Compliance", "Response", "Evidence"],
                ["Yes", "Documented", "Attached"],
            ],
        },
    )
    output_path = tmp_path / "output.json"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path, target="vault_json")
    
    assert exit_code == 0
    
    payload = parse_json_output(stdout)
    assert payload["sheetsProcessed"] == 3
    
    with open(output_path) as f:
        data = json.load(f)
    
    assert data["sheets"]["Compliance"]["type"] == "compliance"
    assert data["sheets"]["Data"]["type"] == "generic"
    assert data["sheets"]["Audit"]["type"] == "compliance"


def test_xlsm_file_support(tmp_path: Path):
    """Verify .xlsm files are supported."""
    # Create .xlsm file (macro-enabled workbook)
    xlsx_path = tmp_path / "test.xlsm"
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is not None:
        ws.title = "Sheet1"
        ws.append(["Name", "Age"])
        ws.append(["Alice", "30"])
    wb.save(str(xlsx_path))
    wb.close()
    
    output_path = tmp_path / "output.md"
    
    exit_code, stdout, stderr = run_converter(xlsx_path, output_path)
    
    assert exit_code == 0
    assert output_path.exists()


if __name__ == "__main__":
    pytest.main([__file__, "-v"])

# Made with Bob
